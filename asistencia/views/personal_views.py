# asistencia/views/personal_views.py
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q, ProtectedError
from django.db import IntegrityError
from django.http import HttpResponse, JsonResponse

from ..models import Personal
from ..forms import PersonalForm


def _filtrar_personal(request):
    query = request.GET.get('q', '').strip()
    filtro = request.GET.get('filtro', '')
    personal = Personal.objects.select_related('cargo', 'institucion', 'distrito', 'nivel', 'turno').all().order_by('-id')
    
    if query:
        personal = personal.filter(
            Q(cod_modular__icontains=query) |
            Q(dni__icontains=query) |
            Q(apellidos__icontains=query) |
            Q(nombres__icontains=query)
        )
    if filtro == 'activos':
        personal = personal.filter(condicion=True)
    elif filtro == 'inactivos':
        personal = personal.filter(condicion=False)
        
    return personal


def personal_listar(request):
    personal_list = _filtrar_personal(request)
    paginator = Paginator(personal_list, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    
    context = {
        'page_obj': page_obj,
        'titulo_seccion': 'Tabla Personal',
        'nombre_singular': 'Personal',
        'campo_busqueda': 'código modular, DNI o apellidos',
        'url_crear': 'personal_crear',
        'url_pdf': 'personal_pdf',
        'url_excel': 'personal_excel',
        'mantenimiento_activo': False,
    }
    return render(request, 'personal/listar.html', context)


def personal_crear(request):
    if request.method == 'POST':
        form = PersonalForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'El personal se registró correctamente.')
            return redirect('personal_listar')
    else:
        form = PersonalForm()
        
    context = {
        'form': form,
        'titulo_formulario': 'Registrar Nuevo Personal',
        'mantenimiento_activo': True,
    }
    return render(request, 'personal/form.html', context)


def personal_editar(request, pk):
    persona = get_object_or_404(Personal, pk=pk)
    if request.method == 'POST':
        form = PersonalForm(request.POST, instance=persona)
        if form.is_valid():
            form.save()
            messages.success(request, f'Los datos de {persona.nombres} {persona.apellidos} fueron actualizados.')
            return redirect('personal_listar')
    else:
        form = PersonalForm(instance=persona)
        
    context = {
        'form': form,
        'persona': persona,
        'titulo_formulario': f'Editar Personal: {persona.nombres} {persona.apellidos}',
        'mantenimiento_activo': True,
    }
    return render(request, 'personal/form.html', context)


def personal_anular(request, pk):
    if request.method == 'POST':
        persona = get_object_or_404(Personal, pk=pk)
        try:
            nombre = f"{persona.nombres} {persona.apellidos}"
            persona.delete()
            messages.success(request, f'El registro de "{nombre}" fue eliminado permanentemente.')
        except (ProtectedError, IntegrityError):
            persona.condicion = False
            persona.save()
            messages.warning(request, f'El personal "{persona.nombres}" tiene registros asociados. Se cambió su condición a Inactivo.')
    return redirect('personal_listar')


# Validaciones en tiempo real
def verificar_unicidad_personal(request):
    campo = request.GET.get('campo', '')
    valor = request.GET.get('valor', '').strip()
    persona_id = request.GET.get('id', None)
    
    exists = False
    if campo and valor:
        filtros = {f"{campo}": valor}
        qs = Personal.objects.filter(**filtros)
        if persona_id and persona_id.isdigit():
            qs = qs.exclude(id=int(persona_id))
        exists = qs.exists()
        
    return JsonResponse({'existe': exists})


# Reporte Excel
def personal_exportar_excel(request):
    personal_list = _filtrar_personal(request)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Personal"
    
    headers = ['ID', 'CÓD. MODULAR', 'DNI', 'APELLIDOS Y NOMBRES', 'CARGO', 'INSTITUCIÓN', 'H. ENTRADA', 'H. SALIDA', 'ESTADO']
    ws.append(headers)
    
    fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for p in personal_list:
        ws.append([
            p.id,
            p.cod_modular,
            p.dni,
            f"{p.apellidos}, {p.nombres}",
            p.cargo.nombre if p.cargo else '',
            p.institucion.nombre if p.institucion else '',
            p.h_entrada.strftime('%H:%M'),
            p.h_salida.strftime('%H:%M'),
            'ACTIVO' if p.condicion else 'INACTIVO'
        ])
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Reporte_Personal.xlsx"'
    wb.save(response)
    return response


# Reporte PDF (ReportLab)
def personal_reporte_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="reporte_personal.pdf"'

    doc = SimpleDocTemplate(response, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    elements = []
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'TitleStyle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=15, alignment=1, textColor=colors.HexColor('#1F4E78')
    )
    elements.append(Paragraph("REPORTE GENERAL DE PERSONAL", title_style))
    elements.append(Spacer(1, 15))

    data = [['ID', 'DNI', 'Cód. Mod.', 'Apellidos y Nombres', 'Cargo', 'Institución', 'Estado']]
    personal_list = _filtrar_personal(request)
    cell_style = ParagraphStyle('CellStyle', fontName='Helvetica', fontSize=8, leading=10)

    for p in personal_list:
        data.append([
            str(p.id),
            p.dni,
            p.cod_modular,
            Paragraph(f"{p.apellidos}, {p.nombres}", cell_style),
            Paragraph(p.cargo.nombre if p.cargo else '', cell_style),
            Paragraph(p.institucion.nombre if p.institucion else '', cell_style),
            'ACTIVO' if p.condicion else 'INACTIVO'
        ])

    t = Table(data, colWidths=[25, 55, 65, 140, 100, 115, 50])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (3, 1), (5, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))

    elements.append(t)
    doc.build(elements)
    return response
