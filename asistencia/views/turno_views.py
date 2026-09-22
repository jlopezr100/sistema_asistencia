import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.core.paginator import Paginator
from django.db.models import ProtectedError
from django.db import IntegrityError

from ..models import Turno
from ..forms import TurnoForm


def _filtrar_turnos(request):
    query = request.GET.get('q', '')
    filtro = request.GET.get('filtro', '')
    turnos = Turno.objects.all().order_by('-id')

    if query:
        turnos = turnos.filter(nombre__icontains=query)
    if filtro == 'activos':
        turnos = turnos.filter(estado=True)
    elif filtro == 'inactivos':
        turnos = turnos.filter(estado=False)
        
    return turnos


def turno_listar(request):
    turnos = _filtrar_turnos(request)
    paginator = Paginator(turnos, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    context = {
        'page_obj': page_obj,
        'titulo_seccion': 'Tabla Turnos',
        'nombre_singular': 'Turno',
        'campo_busqueda': 'nombre de turno',
        'url_crear': 'turno_crear',
        'url_pdf': 'turno_pdf',
        'url_excel': 'turno_excel',
        'mantenimiento_activo': True,
    }
    return render(request, 'turnos/listar.html', context)


def turno_crear(request):
    if request.method == 'POST':
        form = TurnoForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('turno_listar')
    else:
        form = TurnoForm()

    context = {
        'form': form,
        'titulo_formulario': 'Registrar Nuevo Turno',
        'url_cancelar': 'turno_listar',
        'mantenimiento_activo': True,
    }
    return render(request, 'turnos/form.html', context)


def turno_editar(request, pk):
    turno = get_object_or_404(Turno, pk=pk)
    if request.method == 'POST':
        form = TurnoForm(request.POST, instance=turno)
        if form.is_valid():
            form.save()
            return redirect('turno_listar')
    else:
        form = TurnoForm(instance=turno)

    context = {
        'form': form,
        'titulo_formulario': f'Editar Turno: {turno.nombre}',
        'url_cancelar': 'turno_listar',
        'mantenimiento_activo': True,
    }
    return render(request, 'turnos/form.html', context)


def turno_anular(request, pk):
    if request.method == 'POST':
        turno = get_object_or_404(Turno, pk=pk)
        try:
            turno.delete()
        except (ProtectedError, IntegrityError):
            turno.estado = False
            turno.save()
            
    return redirect('turno_listar')


def turno_reporte_pdf(request):
    turnos = _filtrar_turnos(request)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="reporte_turnos.pdf"'

    doc = SimpleDocTemplate(
        response,
        pagesize=letter,
        rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30
    )
    story = []
    styles = getSampleStyleSheet()

    titulo_style = ParagraphStyle(
        'TituloReporte',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#002B49'),
        alignment=1,
        spaceAfter=15
    )
    story.append(Paragraph("REPORTE DE TURNOS - UGEL YUNGAY", titulo_style))
    story.append(Spacer(1, 10))

    data = [['ID', 'Nombre del Turno', 'Hora Entrada', 'Hora Salida', 'Estado']]
    cell_style = ParagraphStyle('Cell', parent=styles['Normal'], fontSize=9, leading=11)

    for t in turnos:
        h_entrada = t.hora_entrada.strftime('%H:%M') if t.hora_entrada else '-'
        h_salida = t.hora_salida.strftime('%H:%M') if t.hora_salida else '-'
        estado_texto = "Activo" if t.estado else "Inactivo"
        data.append([
            str(t.id),
            Paragraph(t.nombre or '', cell_style),
            h_entrada,
            h_salida,
            estado_texto
        ])

    tabla = Table(data, colWidths=[40, 200, 110, 110, 80])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (2, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))

    story.append(tabla)
    doc.build(story)
    return response


def turno_exportar_excel(request):
    turnos = _filtrar_turnos(request)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Turnos"

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
    border_thin = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    ws.merge_cells('A1:E1')
    ws['A1'] = "REPORTE DE TURNOS DE TRABAJO - UGEL YUNGAY"
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='1B365D')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    headers = ['ID', 'Nombre del Turno', 'Hora Entrada', 'Hora Salida', 'Estado']
    ws.append([])
    ws.append(headers)

    ws.row_dimensions[3].height = 24
    for col_num in range(1, 6):
        cell = ws.cell(row=3, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for t in turnos:
        h_entrada = t.hora_entrada.strftime('%H:%M') if t.hora_entrada else '-'
        h_salida = t.hora_salida.strftime('%H:%M') if t.hora_salida else '-'
        estado_txt = "Activo" if t.estado else "Inactivo"
        row_data = [t.id, t.nombre, h_entrada, h_salida, estado_txt]
        ws.append(row_data)

    col_widths = {'A': 10, 'B': 30, 'C': 20, 'D': 20, 'E': 15}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=5):
        for cell in row:
            cell.border = border_thin
            if cell.column in [1, 3, 4, 5]:
                cell.alignment = Alignment(horizontal='center', vertical='center')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="reporte_turnos.xlsx"'
    wb.save(response)
    return response
