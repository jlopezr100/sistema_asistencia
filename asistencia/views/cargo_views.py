import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from django.http import HttpResponse

from django.shortcuts import render, redirect,get_object_or_404
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import ProtectedError
from django.db import IntegrityError
from ..models import Cargo
from ..forms import CargoForm

def cargo_listar(request):
    cargos = _filtrar_cargos(request)
    paginator = Paginator(cargos, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    context = {
        'page_obj': page_obj,
        'titulo_seccion': 'Tabla Cargos',
        'nombre_singular': 'Cargo',
        'campo_busqueda': 'nombre de cargo',
        'url_crear': 'cargo_crear',    # <-- AÑADIR
        'url_pdf': 'cargo_pdf',        # <-- AÑADIR
        'url_excel': 'cargo_excel',    # <-- AÑADIR
        'mantenimiento_activo': True,
    }
    return render(request, 'cargos/listar.html', context)

def cargo_crear(request):
    if request.method == 'POST':
        form = CargoForm(request.POST)
        if form.is_valid():
            form.save()
            #messages.success(request, 'El cargo se registró correctamente.')
            return redirect('cargo_listar')
    else:
        form = CargoForm()

    context = {
        'form': form,
        'mantenimiento_activo': True,
    }
    return render(request, 'cargos/crear.html', context)

def cargo_editar(request, pk):
    cargo = get_object_or_404(Cargo, pk=pk)
    if request.method == 'POST':
        form = CargoForm(request.POST, instance=cargo)
        if form.is_valid():
            form.save()
            #messages.success(request, f'El cargo "{cargo.nombre}" fue actualizado correctamente.')
            return redirect('cargo_listar')
    else:
        form = CargoForm(instance=cargo)

    context = {
        'form': form,
        'titulo_formulario': f'Editar Cargo: {cargo.nombre}',
        'mantenimiento_activo': True,
    }
    return render(request, 'cargos/form.html', context)


def cargo_anular(request, pk):
    if request.method == 'POST':
        cargo = get_object_or_404(Cargo, pk=pk)
        try:
            # Intentamos la eliminación física en BD
            cargo_nombre = cargo.nombre
            cargo.delete()
            #messages.success(request, f'El cargo "{cargo_nombre}" fue eliminado permanentemente.')
        except (ProtectedError, IntegrityError):
            # Si está vinculado con empleados u otra tabla, realizamos la baja lógica (inactivación)
            cargo.estado = False
            cargo.save()
            messages.warning(request, f'El cargo "{cargo.nombre}" está vinculado a otros registros. Se ha cambiado su estado a Inactivo.')
            
    return redirect('cargo_listar')

def _filtrar_cargos(request):
    """Función auxiliar para aplicar los filtros de búsqueda."""
    query = request.GET.get('q', '')
    filtro = request.GET.get('filtro', '')
    cargos = Cargo.objects.all().order_by('-id')

    if query:
        cargos = cargos.filter(nombre__icontains=query)
    if filtro == 'activos':
        cargos = cargos.filter(estado=True)
    elif filtro == 'inactivos':
        cargos = cargos.filter(estado=False)
        
    return cargos


def cargo_reporte_pdf(request):
    cargos = _filtrar_cargos(request)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="reporte_cargos.pdf"'

    doc = SimpleDocTemplate(
        response,
        pagesize=letter,
        rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30
    )
    story = []
    styles = getSampleStyleSheet()

    # Título del reporte
    titulo_style = ParagraphStyle(
        'TituloReporte',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#002B49'),
        alignment=1,
        spaceAfter=15
    )
    story.append(Paragraph("REPORTE DE CARGOS - UGEL YUNGAY", titulo_style))
    story.append(Spacer(1, 10))

    # Encabezados y Datos
    data = [['ID', 'Nombre del Cargo', 'Descripción', 'Estado']]
    cell_style = ParagraphStyle('Cell', parent=styles['Normal'], fontSize=9, leading=11)

    for c in cargos:
        estado_texto = "Activo" if c.estado else "Inactivo"
        data.append([
            str(c.id),
            Paragraph(c.nombre or '', cell_style),
            Paragraph(c.descripcion or '-', cell_style),
            estado_texto
        ])

    # Tabla con anchos de columna fijos
    tabla = Table(data, colWidths=[40, 150, 280, 70])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (3, 0), (3, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))

    story.append(tabla)
    doc.build(story)
    return response


def cargo_exportar_excel(request):
    cargos = _filtrar_cargos(request)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cargos"

    # Estilos para Excel
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
    border_thin = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    # Título principal
    ws.merge_cells('A1:D1')
    ws['A1'] = "REPORTE DE CARGOS - UGEL YUNGAY"
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='1B365D')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    # Encabezados de tabla
    headers = ['ID', 'Nombre del Cargo', 'Descripción', 'Estado']
    ws.append([]) # Línea vacía
    ws.append(headers)

    ws.row_dimensions[3].height = 24
    for col_num in range(1, 5):
        cell = ws.cell(row=3, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Filas de datos
    for c in cargos:
        estado_txt = "Activo" if c.estado else "Inactivo"
        row_data = [c.id, c.nombre, c.descripcion or '-', estado_txt]
        ws.append(row_data)

    # Formato de celda y ancho de columnas
    col_widths = {'A': 10, 'B': 30, 'C': 50, 'D': 15}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = border_thin
            if cell.column in [1, 4]:
                cell.alignment = Alignment(horizontal='center', vertical='center')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="reporte_cargos.xlsx"'
    wb.save(response)
    return response
