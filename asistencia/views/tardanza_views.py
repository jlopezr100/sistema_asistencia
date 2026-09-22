import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.core.paginator import Paginator
from django.db.models import ProtectedError
from django.db import IntegrityError

from ..models import Tardanza
from ..forms import TardanzaForm

def _filtrar_tardanzas(request):
    query = request.GET.get('q', '')
    filtro = request.GET.get('filtro', '')
    tardanzas = Tardanza.objects.all().order_by('-id')

    if query:
        tardanzas = tardanzas.filter(tipo_tardanza__icontains=query)
    if filtro == 'activos':
        tardanzas = tardanzas.filter(estado=True)
    elif filtro == 'inactivos':
        tardanzas = tardanzas.filter(estado=False)
    return tardanzas

def tardanza_listar(request):
    tardanzas = _filtrar_tardanzas(request)
    paginator = Paginator(tardanzas, 10)
    page_obj = paginator.get_page(request.GET.get('page'))

    context = {
        'page_obj': page_obj,
        'titulo_seccion': 'Tabla Tardanzas',
        'nombre_singular': 'Tardanza',
        'campo_busqueda': 'tipo de tardanza',
        'url_crear': 'tardanza_crear',
        'url_pdf': 'tardanza_pdf',
        'url_excel': 'tardanza_excel',
        'mantenimiento_activo': True,
    }
    return render(request, 'tardanzas/listar.html', context)

def tardanza_crear(request):
    if request.method == 'POST':
        form = TardanzaForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('tardanza_listar')
    else:
        form = TardanzaForm()

    context = {
        'form': form,
        'titulo_formulario': 'Registrar Tipo de Tardanza',
        'url_cancelar': 'tardanza_listar',
        'mantenimiento_activo': True,
    }
    return render(request, 'tardanzas/form.html', context)

def tardanza_editar(request, pk):
    tardanza = get_object_or_404(Tardanza, pk=pk)
    if request.method == 'POST':
        form = TardanzaForm(request.POST, instance=tardanza)
        if form.is_valid():
            form.save()
            return redirect('tardanza_listar')
    else:
        form = TardanzaForm(instance=tardanza)

    context = {
        'form': form,
        'titulo_formulario': f'Editar Tardanza: {tardanza.tipo_tardanza}',
        'url_cancelar': 'tardanza_listar',
        'mantenimiento_activo': True,
    }
    return render(request, 'tardanzas/form.html', context)

def tardanza_anular(request, pk):
    if request.method == 'POST':
        tardanza = get_object_or_404(Tardanza, pk=pk)
        try:
            tardanza.delete()
        except (ProtectedError, IntegrityError):
            tardanza.estado = False
            tardanza.save()
    return redirect('tardanza_listar')

def tardanza_reporte_pdf(request):
    tardanzas = _filtrar_tardanzas(request)
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="reporte_tardanzas.pdf"'

    doc = SimpleDocTemplate(response, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    story = []
    styles = getSampleStyleSheet()

    titulo_style = ParagraphStyle('Titulo', parent=styles['Heading1'], fontSize=16, leading=20, textColor=colors.HexColor('#002B49'), alignment=1)
    story.append(Paragraph("REPORTE DE TIPOS DE TARDANZA - UGEL YUNGAY", titulo_style))
    story.append(Spacer(1, 15))

    data = [['ID', 'Tipo de Tardanza', 'Justificable', 'Estado']]
    cell_style = ParagraphStyle('Cell', parent=styles['Normal'], fontSize=9, leading=11)

    for t in tardanzas:
        data.append([
            str(t.id),
            Paragraph(t.tipo_tardanza or '', cell_style),
            "Sí" if t.justificable else "No",
            "Activo" if t.estado else "Inactivo"
        ])

    tabla = Table(data, colWidths=[50, 280, 100, 100])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (2, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(tabla)
    doc.build(story)
    return response

def tardanza_exportar_excel(request):
    tardanzas = _filtrar_tardanzas(request)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tardanzas"

    ws.merge_cells('A1:D1')
    ws['A1'] = "REPORTE DE TIPOS DE TARDANZA - UGEL YUNGAY"
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='1B365D')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    headers = ['ID', 'Tipo de Tardanza', 'Descripción', '¿Justificable?', 'Estado']
    ws.append([])
    ws.append(headers)

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')

    for col in range(1, 6):
        cell = ws.cell(row=3, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for t in tardanzas:
        ws.append([t.id, t.tipo_tardanza, t.descripcion or '', "Sí" if t.justificable else "No", "Activo" if t.estado else "Inactivo"])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="reporte_tardanzas.xlsx"'
    wb.save(response)
    return response
