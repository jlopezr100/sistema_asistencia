import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from django.http import HttpResponse, JsonResponse

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import ProtectedError, Q
from django.db import IntegrityError

from ..models import Institucion
from ..forms import InstitucionForm


def verificar_codigo_modular(request):
    codigo = request.GET.get('codigo', '').strip()
    inst_id = request.GET.get('id', None)
    
    exists = Institucion.objects.filter(codigo_modular=codigo)
    if inst_id and inst_id.isdigit():
        exists = exists.exclude(id=int(inst_id))
        
    return JsonResponse({'existe': exists.exists()})

def _filtrar_instituciones(request):
    query = request.GET.get('q', '').strip()
    filtro = request.GET.get('filtro', '')
    instituciones = Institucion.objects.select_related('turno', 'distrito').all().order_by('-id')
    
    if query:
        # Búsqueda por Código Modular, Nombre o Distrito
        instituciones = instituciones.filter(
            Q(codigo_modular__icontains=query) |
            Q(nombre__icontains=query) |
            Q(distrito__nombre__icontains=query)
        )
    if filtro == 'activos':
        instituciones = instituciones.filter(estado=True)
    elif filtro == 'inactivos':
        instituciones = instituciones.filter(estado=False)
        
    return instituciones


def institucion_listar(request):
    instituciones = _filtrar_instituciones(request)
    paginator = Paginator(instituciones, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    
    context = {
        'page_obj': page_obj,
        'titulo_seccion': 'Tabla Instituciones',
        'nombre_singular': 'Institución',
        'campo_busqueda': 'código modular, nombre o distrito',
        'url_crear': 'institucion_crear',
        'url_pdf': 'institucion_pdf',
        'url_excel': 'institucion_excel',
        'mantenimiento_activo': False,
    }
    return render(request, 'instituciones/listar.html', context)


def institucion_crear(request):
    if request.method == 'POST':
        form = InstitucionForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'La Institución se registró correctamente.')
            return redirect('institucion_listar')
    else:
        form = InstitucionForm()
        
    context = {
        'form': form,
        'titulo_formulario': 'Registrar Nueva Institución Educativa',
        'mantenimiento_activo': True,
    }
    return render(request, 'instituciones/form.html', context)

def institucion_editar(request, pk):
    institucion = get_object_or_404(Institucion, pk=pk)
    if request.method == 'POST':
        form = InstitucionForm(request.POST, instance=institucion)
        if form.is_valid():
            form.save()
            messages.success(request, f'La Institución "{institucion.nombre}" fue actualizada.')
            return redirect('institucion_listar')
    else:
        form = InstitucionForm(instance=institucion)
        
    context = {
        'form': form,
        'institucion': institucion,
        'titulo_formulario': f'Editar Institución: {institucion.nombre}',
        'mantenimiento_activo': True,
    }
    return render(request, 'instituciones/form.html', context)


def institucion_anular(request, pk):
    if request.method == 'POST':
        institucion = get_object_or_404(Institucion, pk=pk)
        try:
            nombre = institucion.nombre
            institucion.delete()
            messages.success(request, f'La institución "{nombre}" fue eliminada permanentemente.')
        except (ProtectedError, IntegrityError):
            institucion.estado = False
            institucion.save()
            messages.warning(request, f'La institución "{institucion.nombre}" está vinculada. Se cambió su estado a Inactivo.')
    return redirect('institucion_listar')


# --- REPORTE EXCEL ---
def institucion_exportar_excel(request):
    instituciones = _filtrar_instituciones(request)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Instituciones"
    
    headers = ['ID', 'CÓDIGO MODULAR', 'NOMBRE DE LA I.E.', 'DISTRITO', 'TURNO', 'ESTADO']
    ws.append(headers)
    
    fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for inst in instituciones:
        ws.append([
            inst.id,
            inst.codigo_modular,
            inst.nombre,
            inst.distrito.nombre if inst.distrito else '',
            inst.turno.nombre if inst.turno else '',
            'ACTIVO' if inst.estado else 'INACTIVO'
        ])
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Reporte_Instituciones.xlsx"'
    wb.save(response)
    return response


# --- REPORTE PDF CON REPORTLAB ---
def institucion_reporte_pdf(request):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="reporte_instituciones.pdf"'

    doc = SimpleDocTemplate(
        response,
        pagesize=letter,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )
    elements = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        alignment=1, # Centrado
        textColor=colors.HexColor('#1F4E78')
    )
    
    elements.append(Paragraph("REPORTE DE INSTITUCIONES EDUCATIVAS", title_style))
    elements.append(Spacer(1, 15))

    # Encabezados de tabla
    data = [['ID', 'Cód. Modular', 'Nombre de la I.E.', 'Distrito', 'Turno', 'Estado']]

    instituciones = _filtrar_instituciones(request)
    
    cell_style = ParagraphStyle('CellStyle', fontName='Helvetica', fontSize=9, leading=11)

    for inst in instituciones:
        estado_str = 'ACTIVO' if inst.estado else 'INACTIVO'
        data.append([
            str(inst.id),
            inst.codigo_modular,
            Paragraph(inst.nombre, cell_style),
            inst.distrito.nombre if inst.distrito else '',
            inst.turno.nombre if inst.turno else '',
            estado_str
        ])

    # Anchos de columna en puntos (Suma = 550 pt aprox)
    t = Table(data, colWidths=[30, 75, 185, 100, 90, 70])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ALIGN', (2, 1), (2, -1), 'LEFT'), # Nombre alineado a la izquierda
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F8F9FA')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))

    elements.append(t)
    doc.build(elements)
    
    return response