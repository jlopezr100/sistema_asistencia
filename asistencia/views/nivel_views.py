import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from django.http import HttpResponse

from django.shortcuts import render, redirect,get_object_or_404
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import ProtectedError
from django.db import IntegrityError
from ..models import Nivel
from ..forms import NivelForm

def _filtrar_niveles(request):
    """Función auxiliar para filtrar niveles según búsqueda y estado."""
    query = request.GET.get('q', '')
    filtro = request.GET.get('filtro', '')
    niveles = Nivel.objects.all().order_by('-id')

    if query:
        niveles = niveles.filter(nombre__icontains=query)
    if filtro == 'activos':
        niveles = niveles.filter(estado=True)
    elif filtro == 'inactivos':
        niveles = niveles.filter(estado=False)
        
    return niveles


def nivel_listar(request):
    niveles = _filtrar_niveles(request)
    paginator = Paginator(niveles, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    context = {
        'page_obj': page_obj,
        'titulo_seccion': 'Tabla Niveles',
        'nombre_singular': 'Nivel',
        'campo_busqueda': 'nombre de nivel',
        'url_crear': 'nivel_crear',
        'url_pdf': 'nivel_pdf',
        'url_excel': 'nivel_excel',
        'mantenimiento_activo': True,
    }
    return render(request, 'niveles/listar.html', context)


def nivel_crear(request):
    if request.method == 'POST':
        form = NivelForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('nivel_listar')
    else:
        form = NivelForm()

    context = {
        'form': form,
        'titulo_formulario': 'Registrar Nuevo Nivel',
        'url_cancelar': 'nivel_listar',
        'mantenimiento_activo': True,
    }
    return render(request, 'niveles/form.html', context)


def nivel_editar(request, pk):
    nivel = get_object_or_404(Nivel, pk=pk)
    if request.method == 'POST':
        form = NivelForm(request.POST, instance=nivel)
        if form.is_valid():
            form.save()
            return redirect('nivel_listar')
    else:
        form = NivelForm(instance=nivel)

    context = {
        'form': form,
        'titulo_formulario': f'Editar Nivel: {nivel.nombre}',
        'url_cancelar': 'nivel_listar',
        'mantenimiento_activo': True,
    }
    return render(request, 'niveles/form.html', context)


def nivel_anular(request, pk):
    if request.method == 'POST':
        nivel = get_object_or_404(Nivel, pk=pk)
        try:
            # Intentamos eliminación física si no está vinculado
            nivel.delete()
        except (ProtectedError, IntegrityError):
            # Si tiene relaciones con otras tablas, realizamos baja lógica
            nivel.estado = False
            nivel.save()
            
    return redirect('nivel_listar')


def nivel_reporte_pdf(request):
    niveles = _filtrar_niveles(request)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="reporte_niveles.pdf"'

    doc = SimpleDocTemplate(
        response,
        pagesize=letter,
        rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30
    )
    story = []
    styles = getSampleStyleSheet()

    titulo_style = ParagraphStyle(
        'TituloReporte',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#002B49'),
        alignment=1,
        spaceAfter=15
    )
    story.append(Paragraph("REPORTE DE NIVELES EDUCATIVOS - UGEL YUNGAY", titulo_style))
    story.append(Spacer(1, 10))

    data = [['ID', 'Nombre del Nivel', 'Descripción', 'Estado']]
    cell_style = ParagraphStyle('Cell', parent=styles['Normal'], fontSize=9, leading=11)

    for n in niveles:
        estado_texto = "Activo" if n.estado else "Inactivo"
        data.append([
            str(n.id),
            Paragraph(n.nombre or '', cell_style),
            Paragraph(n.descripcion or '-', cell_style),
            estado_texto
        ])

    tabla = Table(data, colWidths=[40, 150, 280, 70])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (3, 0), (3, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))

    story.append(tabla)
    doc.build(story)
    return response


def nivel_exportar_excel(request):
    niveles = _filtrar_niveles(request)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Niveles"

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
    border_thin = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    ws.merge_cells('A1:D1')
    ws['A1'] = "REPORTE DE NIVELES EDUCATIVOS - UGEL YUNGAY"
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='1B365D')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    headers = ['ID', 'Nombre del Nivel', 'Descripción', 'Estado']
    ws.append([])
    ws.append(headers)

    ws.row_dimensions[3].height = 24
    for col_num in range(1, 5):
        cell = ws.cell(row=3, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for n in niveles:
        estado_txt = "Activo" if n.estado else "Inactivo"
        row_data = [n.id, n.nombre, n.descripcion or '-', estado_txt]
        ws.append(row_data)

    col_widths = {'A': 10, 'B': 30, 'C': 50, 'D': 15}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = border_thin
            if cell.column in [1, 4]:
                cell.alignment = Alignment(horizontal='center', vertical='center')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="reporte_niveles.xlsx"'
    wb.save(response)
    return response
