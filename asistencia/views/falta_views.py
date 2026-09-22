import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.core.paginator import Paginator
from django.db.models import ProtectedError
from django.db import IntegrityError

from ..models import Falta
from ..forms import FaltaForm

def _filtrar_faltas(request):
    query = request.GET.get('q', '')
    filtro = request.GET.get('filtro', '')
    faltas = Falta.objects.all().order_by('-id')

    if query:
        faltas = faltas.filter(tipo_falta__icontains=query)
    if filtro == 'activos':
        faltas = faltas.filter(estado=True)
    elif filtro == 'inactivos':
        faltas = faltas.filter(estado=False)
    return faltas

def falta_listar(request):
    faltas = _filtrar_faltas(request)
    paginator = Paginator(faltas, 10)
    page_obj = paginator.get_page(request.GET.get('page'))

    context = {
        'page_obj': page_obj,
        'titulo_seccion': 'Tabla Faltas',
        'nombre_singular': 'Falta',
        'campo_busqueda': 'tipo de falta',
        'url_crear': 'falta_crear',
        'url_pdf': 'falta_pdf',
        'url_excel': 'falta_excel',
        'mantenimiento_activo': True,
    }
    return render(request, 'faltas/listar.html', context)

def falta_crear(request):
    if request.method == 'POST':
        form = FaltaForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('falta_listar')
    else:
        form = FaltaForm()

    context = {
        'form': form,
        'titulo_formulario': 'Registrar Tipo de Falta',
        'url_cancelar': 'falta_listar',
        'mantenimiento_activo': True,
    }
    return render(request, 'faltas/form.html', context)

def falta_editar(request, pk):
    falta = get_object_or_404(Falta, pk=pk)
    if request.method == 'POST':
        form = FaltaForm(request.POST, instance=falta)
        if form.is_valid():
            form.save()
            return redirect('falta_listar')
    else:
        form = FaltaForm(instance=falta)

    context = {
        'form': form,
        'titulo_formulario': f'Editar Falta: {falta.tipo_falta}',
        'url_cancelar': 'falta_listar',
        'mantenimiento_activo': True,
    }
    return render(request, 'faltas/form.html', context)

def falta_anular(request, pk):
    if request.method == 'POST':
        falta = get_object_or_404(Falta, pk=pk)
        try:
            falta.delete()
        except (ProtectedError, IntegrityError):
            falta.estado = False
            falta.save()
    return redirect('falta_listar')

def falta_reporte_pdf(request):
    faltas = _filtrar_faltas(request)
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="reporte_faltas.pdf"'

    doc = SimpleDocTemplate(response, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    story = []
    styles = getSampleStyleSheet()

    titulo_style = ParagraphStyle('Titulo', parent=styles['Heading1'], fontSize=16, leading=20, textColor=colors.HexColor('#002B49'), alignment=1)
    story.append(Paragraph("REPORTE DE TIPOS DE FALTA - UGEL YUNGAY", titulo_style))
    story.append(Spacer(1, 15))

    data = [['ID', 'Tipo de Falta', 'Justificable', 'Estado']]
    cell_style = ParagraphStyle('Cell', parent=styles['Normal'], fontSize=9, leading=11)

    for f in faltas:
        data.append([
            str(f.id),
            Paragraph(f.tipo_falta or '', cell_style),
            "Sí" if f.justificable else "No",
            "Activo" if f.estado else "Inactivo"
        ])

    tabla = Table(data, colWidths=[50, 280, 100, 100])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (2, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(tabla)
    doc.build(story)
    return response

def falta_exportar_excel(request):
    faltas = _filtrar_faltas(request)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Faltas"

    ws.merge_cells('A1:D1')
    ws['A1'] = "REPORTE DE TIPOS DE FALTA - UGEL YUNGAY"
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='1B365D')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    headers = ['ID', 'Tipo de Falta', 'Descripción', '¿Justificable?', 'Estado']
    ws.append([])
    ws.append(headers)

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')

    for col in range(1, 6):
        cell = ws.cell(row=3, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for f in faltas:
        ws.append([f.id, f.tipo_falta, f.descripcion or '', "Sí" if f.justificable else "No", "Activo" if f.estado else "Inactivo"])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="reporte_faltas.xlsx"'
    wb.save(response)
    return response
