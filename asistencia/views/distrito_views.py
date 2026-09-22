import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.core.paginator import Paginator
from django.db.models import ProtectedError
from django.db import IntegrityError

from ..models import Distrito
from ..forms import DistritoForm


def _filtrar_distritos(request):
    """Función auxiliar para filtrar distritos por búsqueda y estado."""
    query = request.GET.get('q', '')
    filtro = request.GET.get('filtro', '')
    distritos = Distrito.objects.all().order_by('-id')

    if query:
        distritos = distritos.filter(nombre__icontains=query)
    if filtro == 'activos':
        distritos = distritos.filter(estado=True)
    elif filtro == 'inactivos':
        distritos = distritos.filter(estado=False)
        
    return distritos


def distrito_listar(request):
    distritos = _filtrar_distritos(request)
    paginator = Paginator(distritos, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    context = {
        'page_obj': page_obj,
        'titulo_seccion': 'Tabla Distritos',
        'nombre_singular': 'Distrito',
        'campo_busqueda': 'nombre de distrito',
        'url_crear': 'distrito_crear',
        'url_pdf': 'distrito_pdf',
        'url_excel': 'distrito_excel',
        'mantenimiento_activo': True,
    }
    return render(request, 'distritos/listar.html', context)


def distrito_crear(request):
    if request.method == 'POST':
        form = DistritoForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('distrito_listar')
    else:
        form = DistritoForm()

    context = {
        'form': form,
        'titulo_formulario': 'Registrar Nuevo Distrito',
        'url_cancelar': 'distrito_listar',
        'mantenimiento_activo': True,
    }
    return render(request, 'distritos/form.html', context)


def distrito_editar(request, pk):
    distrito = get_object_or_404(Distrito, pk=pk)
    if request.method == 'POST':
        form = DistritoForm(request.POST, instance=distrito)
        if form.is_valid():
            form.save()
            return redirect('distrito_listar')
    else:
        form = DistritoForm(instance=distrito)

    context = {
        'form': form,
        'titulo_formulario': f'Editar Distrito: {distrito.nombre}',
        'url_cancelar': 'distrito_listar',
        'mantenimiento_activo': True,
    }
    return render(request, 'distritos/form.html', context)


def distrito_anular(request, pk):
    if request.method == 'POST':
        distrito = get_object_or_404(Distrito, pk=pk)
        try:
            # Intentamos eliminación física si no está vinculado
            distrito.delete()
        except (ProtectedError, IntegrityError):
            # Baja lógica si tiene dependencias con Instituciones u otras tablas
            distrito.estado = False
            distrito.save()
            
    return redirect('distrito_listar')


def distrito_reporte_pdf(request):
    distritos = _filtrar_distritos(request)

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'inline; filename="reporte_distritos.pdf"'

    doc = SimpleDocTemplate(
        response,
        pagesize=letter,
        rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
    )
    story = []
    styles = getSampleStyleSheet()

    titulo_style = ParagraphStyle(
        'TituloReporte',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#002B49'),
        alignment=1,
        spaceAfter=15
    )
    story.append(Paragraph("REPORTE DE DISTRITOS - UGEL YUNGAY", titulo_style))
    story.append(Spacer(1, 10))

    data = [['ID', 'Nombre del Distrito', 'Estado']]
    cell_style = ParagraphStyle('Cell', parent=styles['Normal'], fontSize=9, leading=11)

    for d in distritos:
        estado_texto = "Activo" if d.estado else "Inactivo"
        data.append([
            str(d.id),
            Paragraph(d.nombre or '', cell_style),
            estado_texto
        ])

    tabla = Table(data, colWidths=[60, 360, 110])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (2, 0), (2, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')]),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))

    story.append(tabla)
    doc.build(story)
    return response


def distrito_exportar_excel(request):
    distritos = _filtrar_distritos(request)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Distritos"

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1B365D', end_color='1B365D', fill_type='solid')
    border_thin = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    ws.merge_cells('A1:C1')
    ws['A1'] = "REPORTE DE DISTRITOS - UGEL YUNGAY"
    ws['A1'].font = Font(name='Calibri', size=14, bold=True, color='1B365D')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    headers = ['ID', 'Nombre del Distrito', 'Estado']
    ws.append([])
    ws.append(headers)

    ws.row_dimensions[3].height = 24
    for col_num in range(1, 4):
        cell = ws.cell(row=3, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for d in distritos:
        estado_txt = "Activo" if d.estado else "Inactivo"
        row_data = [d.id, d.nombre, estado_txt]
        ws.append(row_data)

    col_widths = {'A': 12, 'B': 45, 'C': 18}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=3):
        for cell in row:
            cell.border = border_thin
            if cell.column in [1, 3]:
                cell.alignment = Alignment(horizontal='center', vertical='center')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="reporte_distritos.xlsx"'
    wb.save(response)
    return response

